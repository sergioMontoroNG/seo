from __future__ import annotations

import asyncio, json, os, re, time, uuid
from html import escape
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import httpx
import psycopg
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook, Workbook
from psycopg.rows import dict_row
from app.faq_parser import parse_faq

HOST=os.getenv('HOST','0.0.0.0'); PORT=int(os.getenv('PORT','3000'))
DB=os.getenv('DATABASE_URL','postgresql://seo:change-me@localhost:5432/seo')
DATA=Path(os.getenv('DATA_DIR','./data')); DATA.mkdir(parents=True,exist_ok=True)
API_VERSION=os.getenv('SHOPIFY_API_VERSION','2026-07')
FAQ_WRITES=os.getenv('FAQ_WRITES_ENABLED','true').lower()=='true'
STORES=json.loads(os.getenv('STORES_JSON','{}'))
FIELDS={
 'handle':('Slug recomendado','handle'), 'title':('H1','title'),
 'seoTitle':('Meta Title','seo.title'), 'seoDescription':('Meta Description','seo.description'),
 'descriptionHtml':('Descripción y Características Técnicas','descriptionHtml'),
 'shortDescription':('Descripción Corta','metafield:custom.descripcion_corta'),
 'faq':('FAQ','metafield:newgarden.faq')}
REQUIRED=['URL Actual','Fecha de consulta','Slug actual',*[x[0] for x in FIELDS.values()]]

def db(): return psycopg.connect(DB, row_factory=dict_row)
def init_db():
 with db() as c:
  c.execute('''CREATE TABLE IF NOT EXISTS runs(id text primary key, store_code text not null, filename text not null, status text not null, created_at timestamptz not null, analyzed_at timestamptz, executed_at timestamptz, created_by text, total_rows int default 0, original_path text not null, rollback_of text, validation_errors jsonb not null default '[]'::jsonb)''')
  c.execute("ALTER TABLE runs ADD COLUMN IF NOT EXISTS validation_errors jsonb NOT NULL DEFAULT '[]'::jsonb")
  c.execute('''CREATE TABLE IF NOT EXISTS run_products(id bigserial primary key, run_id text references runs(id) on delete cascade, row_number int, product_id text, current_handle text, recommended_handle text, url_actual text, status text not null, before_json jsonb, proposed_json jsonb, after_json jsonb, changes_json jsonb, errors_json jsonb, included boolean default true)''')
init_db()
app=FastAPI(title='Newgarden SEO Updater',docs_url=None,redoc_url=None)
app.mount('/static', StaticFiles(directory=Path(__file__).parent/'static'), name='static')

@app.on_event('startup')
def recover_interrupted_runs():
 with db() as c:
  c.execute("update run_products set status='ERROR',errors_json='[\"El servidor se reinició durante la ejecución; comprobar Shopify antes de reintentar.\"]'::jsonb where status in ('UPDATING','REQUESTED') and run_id in (select id from runs where status='EXECUTING')")
  c.execute("update runs set status='COMPLETED_WITH_ERRORS' where status='EXECUTING'")

@app.middleware('http')
async def internal_auth(request: Request, call_next):
 user=os.getenv('AUTH_USER',''); password=os.getenv('AUTH_PASSWORD','')
 if user and password and request.url.path not in ['/health']:
  auth=request.headers.get('authorization','')
  import base64
  try: valid=auth.startswith('Basic ') and base64.b64decode(auth[6:]).decode()==f'{user}:{password}'
  except Exception: valid=False
  if not valid: return Response(status_code=401,headers={'WWW-Authenticate':'Basic realm="Newgarden interno"'})
 return await call_next(request)

@app.get('/health')
def health():
 try:
  with db() as c: c.execute('select 1')
  return {'status':'ok','database':'ok'}
 except Exception: return JSONResponse({'status':'degraded','database':'unavailable'},status_code=503)

TOKEN_CACHE={}
def client_credentials_token(code,s):
 # Apps del Dev Dashboard (B2B, Web Corporativa): el token dura 24 h, se renueva con 5 min de margen.
 cached=TOKEN_CACHE.get(code)
 if cached and cached[1]-300>time.time(): return cached[0]
 cid=os.getenv(s.get('client_id_env','')); secret=os.getenv(s.get('client_secret_env',''))
 if not cid or not secret: raise HTTPException(503,f'Faltan las variables "{s.get("client_id_env")}" / "{s.get("client_secret_env")}" (client credentials) para la tienda {code}. Configúralas en .env y recrea el contenedor app.')
 try:
  r=httpx.post(f"https://{s['domain']}/admin/oauth/access_token",json={'grant_type':'client_credentials','client_id':cid,'client_secret':secret},timeout=30)
  token=r.json().get('access_token') if r.status_code==200 else None
  if not token: raise HTTPException(502,f'Shopify no concedió token a la app de la tienda {code} (HTTP {r.status_code}). Revisa el Client ID/secret y que la app esté instalada.')
  TOKEN_CACHE[code]=(token,time.time()+int(r.json().get('expires_in',3600)))
  return token
 except httpx.HTTPError as e: raise HTTPException(502,f'No se pudo contactar con Shopify para renovar el token de {code}.') from e

def store_cfg(code):
 s=STORES.get(code)
 if not s or not s.get('domain'): raise HTTPException(400,'Tienda no configurada en el servidor.')
 if s.get('client_id_env'): return s,client_credentials_token(code,s)
 env=s.get('token_env',''); token=os.getenv(env) if env else None
 if not token: raise HTTPException(503,f'Falta la variable de entorno "{env}" indicada en token_env para la tienda {code}. Configúrala en .env y recrea el contenedor app.')
 return s,token

async def gql(code, token, query, variables=None):
 s=STORES[code]; url=f"https://{s['domain']}/admin/api/{API_VERSION}/graphql.json"
 async with httpx.AsyncClient(timeout=45) as client:
  for attempt in range(6):
   r=await client.post(url,headers={'X-Shopify-Access-Token':token,'Content-Type':'application/json'},json={'query':query,'variables':variables or {}})
   if r.status_code in (401,403): raise HTTPException(502,'Shopify rechazó la autenticación o los permisos.')
   if r.status_code==429:
    if attempt==5: raise RuntimeError('Shopify mantiene el rate limit tras varios reintentos.')
    await asyncio.sleep(float(r.headers.get('Retry-After',min(2**attempt,20)))); continue
   r.raise_for_status(); body=r.json()
   errors=body.get('errors') or []
   throttled=any((e.get('extensions') or {}).get('code') in ('THROTTLED','RATE_LIMITED') for e in errors)
   if throttled and attempt<5:
    await asyncio.sleep(min(2**attempt,20)); continue
   if errors: raise RuntimeError('GraphQL: '+json.dumps(errors,ensure_ascii=False))
   return body['data']

LOOKUP='''query($identifier: ProductIdentifierInput!){ productByIdentifier(identifier:$identifier){ id handle title descriptionHtml seo{title description} short:metafield(namespace:"custom",key:"descripcion_corta"){type value} faq:metafield(namespace:"newgarden",key:"faq"){type value} } }'''
UPDATE='''mutation($product:ProductUpdateInput!){ productUpdate(product:$product){ product{id handle title descriptionHtml seo{title description} short:metafield(namespace:"custom",key:"descripcion_corta"){value} faq:metafield(namespace:"newgarden",key:"faq"){value}} userErrors{field message} } }'''

def val(v):
 if v is None:return ''
 return str(v).strip()
def parse_xlsx(data:bytes):
 try: wb=load_workbook(filename=__import__('io').BytesIO(data),read_only=True,data_only=True)
 except Exception as e: raise HTTPException(400,'El archivo no es un XLSX válido.') from e
 ws=wb.active
 rows=ws.iter_rows(values_only=True)
 header=next(rows,None)
 if not header: raise HTTPException(400,'La hoja está vacía.')
 headers=[val(x) for x in header]
 duplicates=[h for h in set(headers) if h and headers.count(h)>1]
 if duplicates: raise HTTPException(400,'Encabezados duplicados: '+', '.join(duplicates))
 missing=[h for h in REQUIRED if h not in headers]
 if missing: raise HTTPException(400,'Faltan columnas requeridas: '+', '.join(missing))
 ix={h:i for i,h in enumerate(headers)}; result=[]; errors=[]; handles={}
 for rn,row in enumerate(rows,2):
  if not any(x is not None and str(x).strip() for x in row): continue
  d={h:(row[i] if i<len(row) else None) for h,i in ix.items()}
  slug=val(d['Slug actual'])
  if not slug: errors.append(f'Fila {rn}: Slug actual vacío'); continue
  if slug in handles: errors.append(f'Fila {rn}: Slug actual duplicado ({slug})')
  handles[slug]=rn
  rec=val(d['Slug recomendado'])
  result.append({'rowNumber':rn,'urlActual':val(d['URL Actual']),'fechaConsulta':val(d['Fecha de consulta']), 'currentHandle':slug,'recommendedHandle':rec,
   'title':val(d['H1']),'seoTitle':val(d['Meta Title']),'seoDescription':val(d['Meta Description']), 'descriptionHtml':str(d['Descripción y Características Técnicas']) if d['Descripción y Características Técnicas'] is not None else '', 'shortDescription':str(d['Descripción Corta']) if d['Descripción Corta'] is not None else '', 'faqRaw':str(d['FAQ']) if d['FAQ'] is not None else ''})
 return result,errors

def snapshot(p):
 return {'productId':p['id'],'handle':p['handle'],'title':p['title'],'descriptionHtml':p.get('descriptionHtml') or '', 'seoTitle':(p.get('seo') or {}).get('title'),'seoDescription':(p.get('seo') or {}).get('description'),'shortDescription':(p.get('short') or {}).get('value'),'faqRawValue':(p.get('faq') or {}).get('value')}
def normalize(field,v):
 if v is None:return None
 if field in ('title','seoTitle','seoDescription'): return re.sub(r'\r\n?','\n',v).rstrip()
 return v
def changes_for(src,before,duplicate_handles):
 proposed={k:src.get(k) for k in ['recommendedHandle','title','seoTitle','seoDescription','descriptionHtml','shortDescription','faqRaw']}
 out=[]; errs=[]
 for field,sourcekey,target in [('handle','recommendedHandle','handle'),('title','title','title'),('seoTitle','seoTitle','seoTitle'),('seoDescription','seoDescription','seoDescription'),('descriptionHtml','descriptionHtml','descriptionHtml'),('shortDescription','shortDescription','shortDescription'),('faq','faqRaw','faqRawValue')]:
  new=src.get(sourcekey)
  if new is None or not str(new).strip(): continue
  if field=='handle':
   new=new.strip().lower()
   if not re.fullmatch(r'[a-z0-9]+(?:-[a-z0-9]+)*',new): errs.append('Handle recomendado inválido'); continue
   if new in duplicate_handles: errs.append(f'Handle recomendado duplicado: {new}'); continue
  if field=='faq':
   if not FAQ_WRITES: errs.append('Escritura FAQ deshabilitada por configuración del servidor.'); continue
   try:
    new=parse_faq(new)
   except ValueError as exc: errs.append(f'FAQ no válida: {exc}'); continue
  old=before.get(target)
  if field=='faq' and old:
   try: old=json.dumps(json.loads(old),ensure_ascii=False,separators=(',',':'))
   except Exception: pass
  if normalize(field,old)!=normalize(field,new): out.append({'field':field,'before':old,'after':new,'target':target})
 return out,errs,proposed

@app.get('/',response_class=HTMLResponse)
def home():
 stores=''.join(f'<option value="{escape(str(k), quote=True)}">{escape(str(v.get("name",k)))}</option>' for k,v in STORES.items())
 return (Path(__file__).parent/'templates'/'index.html').read_text(encoding='utf-8').replace('{{stores}}',stores)

def row_out(r):
 return {'id':r['id'],'rowNumber':r['row_number'],'productId':r['product_id'],'currentHandle':r['current_handle'],'recommendedHandle':r['recommended_handle'],'urlActual':r['url_actual'],'status':r['status'],'before':r['before_json'],'proposed':r['proposed_json'],'after':r['after_json'],'changes':r['changes_json'] or [],'errors':r['errors_json'] or [],'included':r['included']}
def get_run(run_id):
 with db() as c:
  run=c.execute('select * from runs where id=%s',(run_id,)).fetchone()
  if not run: raise HTTPException(404,'Run no encontrado')
  products=c.execute('select * from run_products where run_id=%s order by row_number',(run_id,)).fetchall()
 return run,[row_out(p) for p in products]
def run_response(r,products):
 summary={'total':len(products),'found':sum(bool(p['productId']) for p in products),'notFound':sum(p['status']=='NOT_FOUND' for p in products),'withChanges':sum(bool(p['changes']) and p['status'] in ('READY','BLOCKED','VERIFIED') for p in products),'noChanges':sum(p['status']=='NO_CHANGES' for p in products),'blocked':sum(p['status']=='BLOCKED' for p in products),'verified':sum(p['status']=='VERIFIED' for p in products),'ready':r['status']=='ANALYZED' and any(p['status']=='READY' and p['included'] for p in products)}
 return {'id':r['id'],'store_code':r['store_code'],'status':r['status'],'rollbackOf':r.get('rollback_of'),'summary':summary,'products':products,'validationErrors':r.get('validation_errors') or []}

@app.post('/api/runs/analyze')
async def analyze(store_code:str=Form(...),file:UploadFile=File(...)):
 if not file.filename.lower().endswith('.xlsx'): raise HTTPException(400,'Solo se aceptan archivos .xlsx')
 raw=await file.read()
 if len(raw)>25*1024*1024: raise HTTPException(413,'El XLSX supera 25 MB')
 source,errors=parse_xlsx(raw)
 rid=str(uuid.uuid4()); path=DATA/f'{rid}.xlsx'; path.write_bytes(raw)
 with db() as c: c.execute('insert into runs(id,store_code,filename,status,created_at,created_by,total_rows,original_path) values(%s,%s,%s,%s,%s,%s,%s,%s)',(rid,store_code,Path(file.filename).name,'ANALYZING',datetime.now(timezone.utc),os.getenv('AUTH_USER','internal'),len(source),str(path)))
 normalized_recs=[p['recommendedHandle'].strip().lower() for p in source if p['recommendedHandle']]
 duplicate={v for v in normalized_recs if normalized_recs.count(v)>1}
 try:
  cfg,token=store_cfg(store_code)
  # Validate metafield definitions up front; a missing/mismatched definition blocks its field only.
  data=await gql(store_code,token,'query{short:metafieldDefinitions(first:1,ownerType:PRODUCT,namespace:"custom",key:"descripcion_corta"){nodes{type{name}}} faq:metafieldDefinitions(first:1,ownerType:PRODUCT,namespace:"newgarden",key:"faq"){nodes{type{name}}}}')
  defs={ 'short': bool(data['short']['nodes'] and data['short']['nodes'][0]['type']['name']=='multi_line_text_field'), 'faq': bool(data['faq']['nodes'] and data['faq']['nodes'][0]['type']['name']=='rich_text_field') }
 except Exception as e:
  with db() as c: c.execute("update runs set status='FAILED' where id=%s",(rid,))
  raise HTTPException(502,'No se pudo validar la conexión Shopify o las definiciones metafield: '+str(e))
 for src in source:
  errs=[]; before=None; product_id=None; changes=[]; state='PENDING'
  try:
   data=await gql(store_code,token,LOOKUP,{'identifier':{'handle':src['currentHandle']}})
   prod=data['productByIdentifier']
   if not prod:
    # El handle del Excel puede haberse renombrado en Shopify: se sigue el redirect que Shopify creó.
    red=(await gql(store_code,token,'query($q:String!){urlRedirects(first:1,query:$q){nodes{path target}}}',{'q':'path:"/products/'+src['currentHandle']+'"'}))['urlRedirects']['nodes']
    m=re.fullmatch(r'/products/([^/?#]+)',red[0]['target']) if red and red[0]['path']=='/products/'+src['currentHandle'] else None
    if m:
     prod=(await gql(store_code,token,LOOKUP,{'identifier':{'handle':m.group(1)}}))['productByIdentifier']
     if prod and src['recommendedHandle'].strip().lower()==src['currentHandle']: src={**src,'recommendedHandle':prod['handle']}
   if not prod: state='NOT_FOUND'; errs.append('No existe producto con Slug actual.')
   else:
    before=snapshot(prod); product_id=before['productId']
    # Definition checks prevent writes to incompatible custom fields.
    changes,errs,proposed=changes_for(src,before,duplicate)
    if any(c['field']=='shortDescription' for c in changes) and not defs['short']:
     changes=[c for c in changes if c['field']!='shortDescription']; errs.append('Definición custom.descripcion_corta ausente o tipo incorrecto')
    if any(c['field']=='faq' for c in changes) and not defs['faq']:
     changes=[c for c in changes if c['field']!='faq']; errs.append('Definición newgarden.faq ausente o tipo incorrecto')
    critical=any('ocupado' in x or 'duplicado' in x or 'inválido' in x for x in errs)
    if not changes and not errs: state='NO_CHANGES'
    elif critical or (errs and not changes): state='BLOCKED'
    elif changes: state='READY'
    else: state='BLOCKED'
    # Handle occupancy check via exact identifier; never use it to resolve this row's product.
    hc=next((c for c in changes if c['field']=='handle'),None)
    if hc:
     occupied=(await gql(store_code,token,LOOKUP,{'identifier':{'handle':hc['after']}}))['productByIdentifier']
     if occupied and occupied['id']!=product_id: changes.remove(hc); errs.append('Handle recomendado ocupado por otro producto Shopify'); state='BLOCKED'
  except Exception as e: state='ERROR'; errs.append(str(e))
  with db() as c:
   c.execute('insert into run_products(run_id,row_number,product_id,current_handle,recommended_handle,url_actual,status,before_json,proposed_json,changes_json,errors_json) values(%s,%s,%s,%s,%s,%s,%s,%s::jsonb,%s::jsonb,%s::jsonb,%s::jsonb)',(rid,src['rowNumber'],product_id,src['currentHandle'],src['recommendedHandle'] or None,src['urlActual'],state,json.dumps(before) if before else None,json.dumps(src),json.dumps(changes),json.dumps(errs)))
 with db() as c:
  final='VALIDATION_FAILED' if errors else 'ANALYZED'
  c.execute("update runs set status=%s,analyzed_at=now(),validation_errors=%s::jsonb where id=%s",(final,json.dumps(errors),rid))
 run,ps=get_run(rid); return run_response(run,ps)

@app.get('/api/runs')
def runs():
 with db() as c: return [dict(x) for x in c.execute('select id,store_code,status,created_at,total_rows from runs order by created_at desc limit 100').fetchall()]
@app.get('/api/runs/{rid}')
def run_detail(rid):
 r,p=get_run(rid); return run_response(r,p)
@app.patch('/api/products/{pid}')
async def include(pid:int,request:Request):
 body=await request.json()
 with db() as c:
  changed=c.execute('update run_products set included=%s where id=%s returning id',(bool(body.get('included')),pid)).fetchone()
 if not changed: raise HTTPException(404,'Producto no encontrado')
 return {'ok':True}

def build_input(p,current=None):
 inp={'id':p['product_id']}
 for ch in p['changes_json'] or []:
  f=ch['field']; v=ch['after']
  if f in ('title','descriptionHtml'): inp[f]=v
  elif f=='handle': inp['handle']=v; inp['redirectNewHandle']=True
  elif f in ('seoTitle','seoDescription'):
   inp.setdefault('seo',{})['title' if f=='seoTitle' else 'description']=v
  elif f=='shortDescription': inp.setdefault('metafields',[]).append({'namespace':'custom','key':'descripcion_corta','type':'multi_line_text_field','value':v})
  elif f=='faq': inp.setdefault('metafields',[]).append({'namespace':'newgarden','key':'faq','type':'rich_text_field','value':v})
 # ProductUpdateInput.seo replaces the whole object: a missing subfield is cleared in Shopify.
 # Always resend the untouched one from the live snapshot.
 if 'seo' in inp and current:
  for key,field in (('title','seoTitle'),('description','seoDescription')):
   if key not in inp['seo'] and current.get(field) is not None: inp['seo'][key]=current[field]
 return inp
@app.post('/api/runs/{rid}/execute')
async def execute(rid):
 run,ps=get_run(rid)
 if run['status']!='ANALYZED': raise HTTPException(409,'Solo se puede ejecutar un análisis válido.')
 if not any(p['status']=='READY' and p['included'] for p in ps): raise HTTPException(409,'No hay productos listos e incluidos para ejecutar.')
 store_cfg(run['store_code'])
 with db() as c:c.execute("update runs set status='EXECUTING' where id=%s",(rid,))
 asyncio.create_task(execute_worker(rid))
 return {'id':rid,'status':'EXECUTING'}

async def execute_worker(rid):
 run,ps=get_run(rid)
 cfg,token=store_cfg(run['store_code'])
 overall='COMPLETED'
 for p in ps:
  if p['status']!='READY' or not p['included']: continue
  try:
   now=(await gql(run['store_code'],token,LOOKUP,{'identifier':{'handle':p['currentHandle']}}))['productByIdentifier']
   # If handle was changed after analysis, resolve only by GID to preserve identity.
   if not now or now['id']!=p['productId']:
    # GID direct lookup
    now=(await gql(run['store_code'],token,'query($id:ID!){node(id:$id){... on Product{id handle title descriptionHtml seo{title description} short:metafield(namespace:"custom",key:"descripcion_corta"){value} faq:metafield(namespace:"newgarden",key:"faq"){value}}}}',{'id':p['productId']}))['node']
   actual=snapshot(now) if now else None
   if not actual or any(normalize(ch['field'],actual.get(ch['target']))!=normalize(ch['field'],ch['before']) for ch in p['changes']):
    with db() as c: c.execute("update run_products set status='STALE_DATA',errors_json=%s::jsonb where id=%s",(json.dumps(['Shopify cambió desde el análisis; se omitió la escritura.']),p['id']))
    overall='COMPLETED_WITH_ERRORS'; continue
   with db() as c:c.execute("update run_products set status='UPDATING' where id=%s",(p['id'],))
   data=await gql(run['store_code'],token,UPDATE,{'product':build_input({'product_id':p['productId'],'changes_json':p['changes']},actual)})
   result=data['productUpdate']; user_errors=result.get('userErrors') or []
   if user_errors: raise RuntimeError(json.dumps(user_errors,ensure_ascii=False))
   # verify by GID: this remains valid after a handle change.
   afternode=(await gql(run['store_code'],token,'query($id:ID!){node(id:$id){... on Product{id handle title descriptionHtml seo{title description} short:metafield(namespace:"custom",key:"descripcion_corta"){value} faq:metafield(namespace:"newgarden",key:"faq"){value}}}}',{'id':p['productId']}))['node']
   after=snapshot(afternode)
   failed=[ch['field'] for ch in p['changes'] if normalize(ch['field'],after.get(ch['target']))!=normalize(ch['field'],ch['after'])]
   changed={ch['field'] for ch in p['changes']}
   failed+=[f'{f} (borrado sin pedirlo)' for f in ('seoTitle','seoDescription') if f not in changed and normalize(f,after.get(f))!=normalize(f,actual.get(f))]
   if failed: status='VERIFICATION_FAILED'; overall='COMPLETED_WITH_ERRORS'; errs=[f'Verificación no coincide: {", ".join(failed)}']
   else: status='VERIFIED'; errs=[]
   with db() as c:c.execute('update run_products set status=%s,after_json=%s::jsonb,errors_json=%s::jsonb where id=%s',(status,json.dumps(after),json.dumps(errs),p['id']))
  except Exception as e:
   overall='COMPLETED_WITH_ERRORS'
   with db() as c:c.execute("update run_products set status='ERROR',errors_json=%s::jsonb where id=%s",(json.dumps([str(e)]),p['id']))
 with db() as c:
  c.execute('update runs set status=%s,executed_at=now() where id=%s',(overall,rid))
  if run['rollback_of'] and overall=='COMPLETED': c.execute("update runs set status='ROLLED_BACK' where id=%s",(run['rollback_of'],))
 return {'id':rid,'status':overall}

@app.post('/api/runs/{rid}/rollback')
async def rollback(rid):
 original,products=get_run(rid)
 if original['status'] not in ('COMPLETED','COMPLETED_WITH_ERRORS','ROLLED_BACK'):
  raise HTTPException(409,'Solo se puede revertir una ejecución que haya terminado.')
 candidates=[p for p in products if p['status']=='VERIFIED' and p['after'] and p['included']]
 if not candidates: raise HTTPException(409,'No hay productos verificados para revertir.')
 new_id=str(uuid.uuid4())
 with db() as c:
  c.execute("insert into runs(id,store_code,filename,status,created_at,created_by,total_rows,original_path,rollback_of) values(%s,%s,%s,'ANALYZED',now(),%s,%s,%s,%s)",(new_id,original['store_code'],f'Rollback de {rid}',os.getenv('AUTH_USER','internal'),len(candidates),original['original_path'],rid))
  for p in candidates:
   reverse=[]
   for ch in p['changes']:
    reverse.append({**ch,'before':ch['after'],'after':ch['before']})
   c.execute('''insert into run_products(run_id,row_number,product_id,current_handle,recommended_handle,url_actual,status,before_json,proposed_json,changes_json,errors_json)
    values(%s,%s,%s,%s,%s,%s,'READY',%s::jsonb,%s::jsonb,%s::jsonb,'[]'::jsonb)''',
    (new_id,p['rowNumber'],p['productId'],p['after']['handle'],p['before']['handle'],p['urlActual'],json.dumps(p['after']),json.dumps({'rollbackOf':rid}),json.dumps(reverse)))
 original2,ps=get_run(new_id); return run_response(original2,ps)

@app.get('/api/runs/{rid}/matrixify')
def matrixify(rid):
 run,ps=get_run(rid)
 wb=Workbook(); ws=wb.active; ws.title='Products'
 eligible=[p for p in ps if p['status']=='READY' and p['included']]
 names={'handle':'Handle','title':'Title','descriptionHtml':'Body HTML','seoTitle':'SEO Title','seoDescription':'SEO Description','shortDescription':'Metafield: custom.descripcion_corta [multi_line_text_field]','faq':'Metafield: newgarden.faq [rich_text_field]'}
 fields=list(dict.fromkeys(ch['field'] for p in eligible for ch in p['changes']))
 headers=['ID','Handle']+[names[f] for f in fields]; ws.append(headers)
 for p in eligible:
  byfield={ch['field']:ch for ch in p['changes']}; before=p['before'] or {}; values=[]
  values.append(p['productId'].rsplit('/',1)[-1])
  values.append(byfield['handle']['after'] if 'handle' in byfield else p['currentHandle'])
  for f in fields:
   if f in byfield: values.append(byfield[f]['after'])
   else:
    target={'handle':'handle','title':'title','descriptionHtml':'descriptionHtml','seoTitle':'seoTitle','seoDescription':'seoDescription','shortDescription':'shortDescription','faq':'faqRawValue'}[f]
    values.append(before.get(target))
  ws.append(values)
 import io
 bio=io.BytesIO();wb.save(bio)
 return Response(bio.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="matrixify-{rid}.xlsx"'})
